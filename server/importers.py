"""Bounded spreadsheet parsing; column indices, not header names, identify fields."""
import csv
import io
import re
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from .common import APIError, CUSTOMER_COLUMNS

MAX_ROWS = 20000
MAX_COLUMNS = 100
MAX_CELLS = 500000


def read_grid(path, sheet=None, encoding='auto', delimiter='auto'):
    path = Path(path)
    if path.suffix.lower() == '.xlsx':
        try:
            with zipfile.ZipFile(path) as archive:
                if sum(entry.file_size for entry in archive.infolist()) > 50 * 1024 * 1024:
                    raise APIError('압축을 푼 엑셀 파일이 너무 큽니다.', 413)
            # Own the file handle too: interrupted read-only row generators can
            # otherwise leave a Windows file locked after Workbook.close().
            workbook_source = path.open('rb')
            workbook = None
            try:
                workbook = load_workbook(workbook_source, read_only=True, data_only=False)
                sheets = workbook.sheetnames
                selected = sheet or next((name for name in sheets if any(
                    any(cell.value is not None for cell in row)
                    for row in list(workbook[name].iter_rows(max_row=100, max_col=MAX_COLUMNS)))), sheets[0])
                if selected not in sheets:
                    raise APIError('시트를 다시 선택해 주세요.')
                grid, kinds, cells = [], [], 0
                for row in workbook[selected].iter_rows():
                    cells += len(row)
                    if len(grid) >= MAX_ROWS + 100 or len(row) > MAX_COLUMNS or cells > MAX_CELLS:
                        raise APIError('파일의 행 또는 열이 너무 많습니다.', 413)
                    grid.append(['' if cell.value is None else str(cell.value) for cell in row])
                    kinds.append([cell.data_type for cell in row])
            finally:
                if workbook is not None:
                    workbook.close()
                workbook_source.close()
            return grid, kinds, sheets, selected
        except APIError:
            raise
        except Exception as error:
            raise APIError('엑셀 파일을 읽을 수 없습니다. .xlsx 파일인지 확인해 주세요.') from error
    if path.suffix.lower() != '.csv':
        raise APIError('.xlsx 또는 .csv 파일을 선택해 주세요.')
    raw = path.read_bytes()
    text = None
    choices = ('utf-8-sig', 'cp949') if encoding == 'auto' else (encoding,)
    if any(choice not in ('utf-8', 'utf-8-sig', 'cp949') for choice in choices):
        raise APIError('지원하지 않는 문자 인코딩입니다.')
    for choice in choices:
        try:
            text = raw.decode(choice)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise APIError('문자가 깨집니다. UTF-8 또는 CP949를 선택해 다시 확인해 주세요.')
    if '\x00' in text:
        raise APIError('CSV 텍스트 파일이 아닙니다.')
    if delimiter == 'auto':
        try:
            delimiter = csv.Sniffer().sniff(text[:8192], delimiters=',;\t').delimiter
        except csv.Error:
            delimiter = ','
    if delimiter not in (',', ';', '\t'):
        raise APIError('구분자는 쉼표, 세미콜론 또는 탭을 선택해 주세요.')
    grid, cells = [], 0
    try:
        for row in csv.reader(io.StringIO(text, newline=''), delimiter=delimiter):
            cells += len(row)
            if len(grid) >= MAX_ROWS + 100 or len(row) > MAX_COLUMNS or cells > MAX_CELLS:
                raise APIError('파일의 행 또는 열이 너무 많습니다.', 413)
            if any(len(value) > 10000 for value in row):
                raise APIError('너무 긴 셀 값이 있습니다.')
            grid.append(row)
    except csv.Error as error:
        raise APIError('CSV 형식을 확인해 주세요.') from error
    return grid, [['s'] * len(row) for row in grid], [], None


def describe(path, settings=None):
    settings = settings or {}
    grid, _, sheets, selected = read_grid(path, settings.get('sheet'),
                                         settings.get('encoding', 'utf-8-sig'), settings.get('delimiter', ','))
    header = settings.get('header_row', 0)
    if header is not None and (type(header) is not int or not 0 <= header < min(len(grid), 100)):
        raise APIError('제목 줄을 다시 선택해 주세요.')
    width = max((len(row) for row in grid), default=0)
    labels = grid[header] if header is not None else []
    start = header + 1 if header is not None else 0
    columns = [{'index': index, 'label': (labels[index] if index < len(labels) else '') or f'열 {index + 1}',
                'examples': [row[index] for row in grid[start:start + 3] if index < len(row)]}
               for index in range(width)]
    return {'sheets': sheets, 'sheet': selected, 'header_row': header,
            'encoding': settings.get('encoding', 'utf-8-sig'), 'delimiter': settings.get('delimiter', ','),
            'columns': columns,
            'sample': grid[:10], 'row_count': max(0, len(grid) - start)}


def mapped_rows(path, settings):
    mapping = settings.get('mapping', {})
    if not isinstance(mapping, dict) or set(mapping) != set(CUSTOMER_COLUMNS):
        raise APIError('이름, 전화번호, 주소 열을 모두 지정해 주세요.')
    indices = list(mapping.values())
    if any(type(index) is not int or index < 0 or index >= MAX_COLUMNS for index in indices) or len(set(indices)) != 3:
        raise APIError('서로 다른 세 열을 선택해 주세요.')
    grid, kinds, _, _ = read_grid(path, settings.get('sheet'), settings.get('encoding', 'auto'),
                                  settings.get('delimiter', 'auto'))
    header = settings.get('header_row', 0)
    if header is not None and (type(header) is not int or not 0 <= header < min(len(grid), 100)):
        raise APIError('제목 줄을 다시 선택해 주세요.')
    start = header + 1 if header is not None else 0
    rows, errors, seen, duplicates = [], [], set(), 0
    for offset, row in enumerate(grid[start:], start):
        if not any(value.strip() for value in row):
            continue
        if len(rows) + len(errors) >= MAX_ROWS:
            raise APIError('고객 데이터는 20,000행까지 가져올 수 있습니다.', 413)
        customer, issues = {}, []
        for field, index in mapping.items():
            value = row[index].strip() if index < len(row) else ''
            kind = kinds[offset][index] if index < len(kinds[offset]) else 's'
            customer[field] = value
            if not value:
                issues.append(field + ': 값이 없습니다')
            elif kind in ('f', 'e'):
                issues.append(field + ': 수식/오류 셀은 값으로 바꿔 주세요')
            elif len(value) > 1000:
                issues.append(field + ': 값이 너무 깁니다')
            elif field == 'ph' and (kind == 'n' or re.fullmatch(r'\d+\.0|\d+(\.\d+)?[eE][+-]?\d+|1[016789]\d{7,8}', value)):
                issues.append('ph: 숫자 셀 또는 전화번호 앞자리 손실 가능성. 텍스트로 확인해 주세요')
        if issues:
            errors.append({'row': offset + 1, 'issues': issues})
            continue
        key = tuple(customer[field] for field in CUSTOMER_COLUMNS)
        if key in seen:
            duplicates += 1
        else:
            seen.add(key)
            rows.append(customer)
    return rows, errors, duplicates
